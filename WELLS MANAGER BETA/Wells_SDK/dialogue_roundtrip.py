# -*- coding: utf-8 -*-
"""Wells dialogue round-trip built on Ren'Py's native dialogue map.

Ren'Py discovers the source structure and writes dialogue.tab. Wells turns that
into a physical JSON reconstruction map plus a clean TXT/DOCX document. The
translated document is then written into Ren'Py's generated game/tl tree.
"""
from __future__ import annotations

import ast
from datetime import datetime
import json
import re
import shutil
from pathlib import Path

from docx import Document

import sdk_core

MAP_VERSION = 1
MAP_NAME = "Wells_Translation_Map.json"
TEMP_DIR_NAME = "Wells_Temp"
TXT_NAME = "Wells_Dialogue.txt"
DOCX_NAME = "Wells_Dialogue.docx"
LOG_NAME = "Wells_Dialogue_Log.txt"

_TRANSLATE = re.compile(r'^\s*translate\s+([A-Za-z0-9_]+)\s+([A-Za-z0-9_]+|strings)\s*:\s*$')
_OLD = re.compile(r'^\s*old\s+(".*")\s*$')
_NEW = re.compile(r'^\s*new\s+(".*")\s*$')
_QUOTED = re.compile(r'"((?:\\.|[^"\\])*)"')
_SOURCE = re.compile(r'^\s*#\s*(.+):(\d+)\s*$')
_RECORD_SERIAL = re.compile(r'^W[0-9A-F]{8}$', re.I)


class DialogueRoundtripError(RuntimeError):
    pass


def _notify(progress, value, text=None):
    if progress:
        progress(max(0.0, min(100.0, float(value))), text)


def _log_file(output_dir, message):
    path = Path(output_dir) / LOG_NAME
    with path.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(str(message) + "\n")


def _decode_quoted(token):
    return ast.literal_eval(token)


def _quote_renpy(text):
    return '"' + (str(text).replace('\\', '\\\\').replace('"', '\\"').replace('\r', '\\r').replace('\n', '\\n')) + '"'


def _extract_last_string(line):
    matches = list(_QUOTED.finditer(line))
    if not matches:
        return None
    m = matches[-1]
    return _decode_quoted(m.group(0)), m.start(), m.end()


def _replace_last_string(line, value):
    found = _extract_last_string(line)
    if not found:
        raise DialogueRoundtripError("Linha Ren'Py sem texto reconhecível.")
    _, start, end = found
    return line[:start] + _quote_renpy(value) + line[end:]


def _norm_file(value):
    return str(value or "").replace('\\', '/').lstrip('./').casefold()


def _unescape_tab_text(value):
    """Undo DialogueFile's TAB escaping without interpreting arbitrary escapes."""
    out = []
    i = 0
    while i < len(value):
        if value[i] != '\\' or i + 1 >= len(value):
            out.append(value[i]); i += 1; continue
        nxt = value[i + 1]
        if nxt == 'n': out.append('\n'); i += 2
        elif nxt == 't': out.append('\t'); i += 2
        elif nxt == '\\': out.append('\\'); i += 2
        else: out.append('\\'); i += 1
    return ''.join(out)


def parse_dialogue_tab(path):
    path = Path(path)
    lines = path.read_text(encoding="utf-8-sig", errors="strict").splitlines()
    if not lines:
        raise DialogueRoundtripError("dialogue.tab está vazio.")
    header = lines[0].split('\t')
    expected = ["Identifier", "Character", "Dialogue", "Filename", "Line Number", "Ren'Py Script"]
    if header[:6] != expected:
        raise DialogueRoundtripError("Cabeçalho de dialogue.tab inesperado.")
    rows = []
    for number, raw in enumerate(lines[1:], 2):
        if not raw:
            continue
        parts = raw.split('\t')
        if len(parts) < 5:
            raise DialogueRoundtripError("Linha {} inválida em dialogue.tab.".format(number))
        parts += [''] * (6 - len(parts))
        try:
            line_number = int(parts[4])
        except ValueError:
            line_number = 0
        rows.append({
            "identifier": parts[0],
            "character": parts[1],
            "dialogue": _unescape_tab_text(parts[2]),
            "filename": parts[3],
            "line_number": line_number,
            "renpy_script": parts[5],
        })
    return rows


def _source_ref(line):
    m = _SOURCE.match(line.rstrip('\r\n'))
    if not m:
        return None
    return m.group(1), int(m.group(2))


def scan_translation_dir(tl_dir):
    """Locate exact generated TL targets while retaining their source metadata."""
    root = Path(tl_dir).resolve()
    entries = []
    for path in sorted(root.rglob('*.rpy')):
        raw = path.read_bytes()
        bom = raw.startswith(b'\xef\xbb\xbf')
        text = raw.decode('utf-8-sig')
        lines = text.splitlines(keepends=True)
        language = None
        block = None
        is_strings = False
        block_source = ("", 0)
        string_source = ("", 0)
        i = 0
        while i < len(lines):
            body = lines[i].rstrip('\r\n')
            sr = _source_ref(body)
            if sr:
                string_source = sr

            tm = _TRANSLATE.match(body)
            if tm:
                language, block = tm.group(1), tm.group(2)
                is_strings = block == 'strings'
                # Dialogue source comment normally sits immediately above translate.
                j = i - 1
                while j >= 0 and not lines[j].strip():
                    j -= 1
                ref = _source_ref(lines[j]) if j >= 0 else None
                block_source = ref or ("", 0)
                i += 1
                continue

            if language and is_strings:
                om = _OLD.match(body)
                if om:
                    try:
                        original = _decode_quoted(om.group(1))
                    except Exception:
                        i += 1; continue
                    j = i + 1
                    while j < len(lines):
                        candidate = lines[j].rstrip('\r\n')
                        if _TRANSLATE.match(candidate) or _OLD.match(candidate):
                            break
                        nm = _NEW.match(candidate)
                        if nm:
                            try:
                                current = _decode_quoted(nm.group(1))
                            except Exception:
                                break
                            entries.append({
                                "kind": "string", "identifier": "", "character": "",
                                "original": original, "current": current,
                                "source_filename": string_source[0], "source_line": string_source[1],
                                "tl_file": str(path.relative_to(root)).replace('\\', '/'),
                                "target_line": j, "bom": bom,
                            })
                            i = j
                            break
                        j += 1

            elif language and not is_strings:
                stripped = body.lstrip()
                if stripped.startswith('#') and not _source_ref(body):
                    comment = stripped[1:].lstrip()
                    original_info = _extract_last_string(comment)
                    if original_info:
                        original, start, _ = original_info
                        character = comment[:start].strip().split()[0] if comment[:start].strip() else ""
                        j = i + 1
                        while j < len(lines):
                            candidate = lines[j].rstrip('\r\n')
                            if _TRANSLATE.match(candidate):
                                break
                            cstrip = candidate.lstrip()
                            if not candidate.strip() or cstrip.startswith('voice '):
                                j += 1; continue
                            if cstrip.startswith('#'):
                                break
                            target = _extract_last_string(candidate)
                            if target:
                                entries.append({
                                    "kind": "dialogue", "identifier": block or "", "character": character,
                                    "original": original, "current": target[0],
                                    "source_filename": block_source[0], "source_line": block_source[1],
                                    "tl_file": str(path.relative_to(root)).replace('\\', '/'),
                                    "target_line": j, "bom": bom,
                                })
                            break
            i += 1
    return root, entries


def _pick_tab_row(entry, rows, used):
    candidates = []
    for idx, row in enumerate(rows):
        if idx in used or row['dialogue'] != entry['original']:
            continue
        if entry['kind'] == 'dialogue' and entry['identifier'] and row['identifier'] == entry['identifier']:
            candidates.append((0 if row['character'] == entry['character'] else 1, idx, row))
            continue
        same_file = _norm_file(row['filename']) == _norm_file(entry['source_filename'])
        same_line = int(row['line_number'] or 0) == int(entry['source_line'] or 0)
        if same_file and same_line:
            candidates.append((2, idx, row))
    if not candidates:
        return None
    _, idx, row = sorted(candidates, key=lambda item: item[0])[0]
    used.add(idx)
    return row


def _protected_segments(text):
    """Yield Ren'Py []/{} segments plus physical newline/tab characters."""
    result = []
    i = 0
    while i < len(text):
        if text[i] in '\n\t':
            result.append((i, i + 1, text[i])); i += 1; continue
        if text[i] not in '[{':
            i += 1; continue
        opener = text[i]
        closer = ']' if opener == '[' else '}'
        depth = 1
        j = i + 1
        while j < len(text) and depth:
            if text[j] == opener: depth += 1
            elif text[j] == closer: depth -= 1
            j += 1
        if depth:
            result.append((i, len(text), text[i:])); break
        result.append((i, j, text[i:j])); i = j
    return result


def _protect_text(text, entry_index):
    segments = _protected_segments(text)
    if not segments:
        return text, []
    out = []
    last = 0
    protected = []
    for token_index, (start, end, value) in enumerate(segments, 1):
        serial = "W9P{:06X}{:02X}".format(entry_index, token_index)
        out.append(text[last:start]); out.append(serial)
        protected.append({"serial": serial, "value": value})
        last = end
    out.append(text[last:])
    return ''.join(out), protected


def _restore_text(text, protected, record_serial):
    result = text
    for item in protected:
        token = item['serial']
        matches = list(re.finditer(re.escape(token), result, flags=re.I))
        if len(matches) != 1:
            raise DialogueRoundtripError(
                "{}: marcador protegido {} deveria aparecer uma vez, mas apareceu {}. Nada foi alterado.".format(
                    record_serial, token, len(matches)))
        m = matches[0]
        result = result[:m.start()] + item['value'] + result[m.end():]
    return result


def _write_txt(path, records):
    with Path(path).open('w', encoding='utf-8', newline='\n') as handle:
        for serial, text in records:
            handle.write(serial + '\n')
            handle.write(text.replace('\r', ' ') + '\n')


def _write_docx(path, records):
    doc = Document()
    for serial, text in records:
        doc.add_paragraph(serial)
        doc.add_paragraph(text)
    doc.save(str(path))


def _read_txt_units(path):
    return Path(path).read_text(encoding='utf-8-sig').splitlines()


def _read_docx_units(path):
    doc = Document(str(path))
    return [p.text for p in doc.paragraphs]


def _parse_document(path, fmt):
    units = _read_docx_units(path) if fmt == 'docx' else _read_txt_units(path)
    found = {}
    i = 0
    while i < len(units):
        candidate = units[i].strip()
        if not _RECORD_SERIAL.fullmatch(candidate):
            i += 1; continue
        serial = candidate.upper()
        if serial in found:
            raise DialogueRoundtripError("Serial duplicado no documento: {}.".format(serial))
        i += 1
        content = []
        while i < len(units) and not _RECORD_SERIAL.fullmatch(units[i].strip()):
            content.append(units[i]); i += 1
        while content and not content[0].strip(): content.pop(0)
        while content and not content[-1].strip(): content.pop()
        # Newlines originating in Ren'Py are protected by W9P... markers, so
        # extra physical wrapping introduced by an editor can safely be joined.
        found[serial] = ' '.join(x.strip() for x in content if x.strip())
    return found


def prepare(project, language, fmt='txt', output_dir=None, include_strings=True, log=None, progress=None):
    if fmt not in ('txt', 'docx'):
        raise DialogueRoundtripError("Formato Wells inválido.")
    project = sdk_core._project_root(project)
    language = sdk_core._language(language, 'Informe o idioma da tradução.')
    output_dir = Path(output_dir or project).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    temp_dir = output_dir / TEMP_DIR_NAME
    temp_dir.mkdir(parents=True, exist_ok=True)
    map_path = temp_dir / MAP_NAME
    tab_copy = temp_dir / 'dialogue.tab'
    document_path = output_dir / (DOCX_NAME if fmt == 'docx' else TXT_NAME)

    _log_file(output_dir, "=== WELLS DIALOGUE PREPARE {} ===".format(datetime.now().isoformat(timespec='seconds')))
    _notify(progress, 3, "Gerando estrutura de tradução Ren'Py...")
    sdk_core.generate_translations(project, language, empty=True, log=log)
    _notify(progress, 20, "Extraindo mapa TAB nativo...")
    tab_result = sdk_core.extract_dialogue(project, fmt='tab', strings=include_strings, notags=False, escape=False, log=log)
    tab_source = Path(tab_result['file'])
    rows = parse_dialogue_tab(tab_source)
    if tab_copy.exists(): tab_copy.unlink()
    shutil.move(str(tab_source), str(tab_copy))

    _notify(progress, 38, "Mapeando destinos game/tl...")
    tl_dir = project / 'game' / 'tl' / language
    root, tl_entries = scan_translation_dir(tl_dir)
    if not tl_entries:
        raise DialogueRoundtripError("Nenhum registro traduzível foi encontrado em {}.".format(tl_dir))

    used = set()
    map_entries = []
    document_records = []
    unmatched = 0
    total = max(1, len(tl_entries))
    for index, entry in enumerate(tl_entries, 1):
        row = _pick_tab_row(entry, rows, used)
        if row is None:
            unmatched += 1
            row = {
                'identifier': entry['identifier'], 'character': entry['character'],
                'dialogue': entry['original'], 'filename': entry['source_filename'],
                'line_number': entry['source_line'], 'renpy_script': ''}
        serial = "W{:08X}".format(index)
        protected_text, protected = _protect_text(entry['original'], index)
        document_records.append((serial, protected_text))
        map_entries.append({
            "serial": serial,
            "kind": entry['kind'],
            "identifier": row['identifier'] or entry['identifier'],
            "character": row['character'] or entry['character'],
            "dialogue": row['dialogue'],
            "filename": row['filename'] or entry['source_filename'],
            "line_number": row['line_number'] or entry['source_line'],
            "renpy_script": row['renpy_script'],
            "tl_file": entry['tl_file'],
            "target_line": entry['target_line'],
            "tl_current": entry['current'],
            "protected": protected,
        })
        if index % 500 == 0 or index == len(tl_entries):
            _notify(progress, 38 + 35 * index / total, "Criando mapa: {}/{}".format(index, len(tl_entries)))

    payload = {
        "wells_map_version": MAP_VERSION,
        "created": datetime.now().isoformat(timespec='seconds'),
        "project": project.name,
        "language": language,
        "format": fmt,
        "document": document_path.name,
        "translation_root": str(root),
        "tab_file": tab_copy.name,
        "entries": map_entries,
    }
    map_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    _notify(progress, 80, "Gravando documento Wells...")
    if fmt == 'docx': _write_docx(document_path, document_records)
    else: _write_txt(document_path, document_records)

    _log_file(output_dir, "Projeto: {}".format(project.name))
    _log_file(output_dir, "Idioma: {}".format(language))
    _log_file(output_dir, "Registros: {}".format(len(map_entries)))
    _log_file(output_dir, "TAB associados: {}; fallback TL: {}".format(len(map_entries) - unmatched, unmatched))
    _log_file(output_dir, "Documento: {}".format(document_path.name))
    _log_file(output_dir, "Mapa físico: {}".format(map_path))
    _notify(progress, 100, "Documento e mapa Wells prontos.")
    return {"document": str(document_path), "map": str(map_path), "entries": len(map_entries), "unmatched": unmatched}


def inject(project, document, log=None, progress=None):
    project = sdk_core._project_root(project)
    document = Path(document).expanduser().resolve()
    if not document.is_file():
        raise DialogueRoundtripError("Documento traduzido não encontrado: {}".format(document))
    output_dir = document.parent
    map_path = output_dir / TEMP_DIR_NAME / MAP_NAME
    if not map_path.is_file():
        raise DialogueRoundtripError("Mapa físico Wells não encontrado ao lado deste documento: {}".format(map_path))
    payload = json.loads(map_path.read_text(encoding='utf-8'))
    if payload.get('wells_map_version') != MAP_VERSION:
        raise DialogueRoundtripError("Versão do mapa Wells incompatível.")
    if payload.get('project') != project.name:
        raise DialogueRoundtripError("Este mapa pertence ao projeto '{}', não a '{}'.".format(payload.get('project'), project.name))
    fmt = payload.get('format')
    if fmt not in ('txt', 'docx'):
        raise DialogueRoundtripError("Formato inválido no mapa Wells.")
    if document.suffix.lower() != '.' + fmt:
        raise DialogueRoundtripError("O mapa espera um documento {}.".format(fmt.upper()))

    _log_file(output_dir, "=== WELLS DIALOGUE INJECT {} ===".format(datetime.now().isoformat(timespec='seconds')))
    _notify(progress, 8, "Lendo documento traduzido...")
    translated = _parse_document(document, fmt)
    entries = payload.get('entries') or []
    expected_serials = [e['serial'].upper() for e in entries]
    missing = [s for s in expected_serials if s not in translated]
    extras = [s for s in translated if s not in set(expected_serials)]
    if missing or extras:
        raise DialogueRoundtripError(
            "Documento não corresponde ao mapa. Ausentes: {}; extras: {}. Nada foi alterado.".format(
                ', '.join(missing[:10]) or '0', ', '.join(extras[:10]) or '0'))

    language = payload['language']
    tl_root = project / 'game' / 'tl' / language
    per_file = {}
    total = max(1, len(entries))
    for index, entry in enumerate(entries, 1):
        serial = entry['serial'].upper()
        value = _restore_text(translated[serial], entry.get('protected') or [], serial)
        path = (tl_root / entry['tl_file']).resolve()
        try:
            path.relative_to(tl_root.resolve())
        except ValueError:
            raise DialogueRoundtripError("Destino fora da pasta TL no mapa: {}".format(path))
        per_file.setdefault(path, []).append((entry, value))
        if index % 500 == 0 or index == len(entries):
            _notify(progress, 8 + 32 * index / total, "Validando documento: {}/{}".format(index, len(entries)))

    prepared = {}
    originals = {}
    files = list(per_file.items())
    for file_index, (path, replacements) in enumerate(files, 1):
        if not path.is_file():
            raise DialogueRoundtripError("Arquivo TL não encontrado: {}. Nada foi alterado.".format(path))
        raw = path.read_bytes(); originals[path] = raw
        bom = raw.startswith(b'\xef\xbb\xbf')
        lines = raw.decode('utf-8-sig').splitlines(keepends=True)
        for entry, value in replacements:
            line_index = int(entry['target_line'])
            if line_index < 0 or line_index >= len(lines):
                raise DialogueRoundtripError("Estrutura TL mudou em {}. Nada foi alterado.".format(entry['tl_file']))
            ending = '\r\n' if lines[line_index].endswith('\r\n') else ('\n' if lines[line_index].endswith('\n') else '')
            body = lines[line_index][:-len(ending)] if ending else lines[line_index]
            current = _extract_last_string(body)
            if not current or current[0] != entry.get('tl_current', ''):
                raise DialogueRoundtripError(
                    "Destino TL mudou em {}:{} desde a extração. Nada foi alterado.".format(entry['tl_file'], line_index + 1))
            lines[line_index] = _replace_last_string(body, value) + ending
        data = ''.join(lines).encode('utf-8')
        if bom: data = b'\xef\xbb\xbf' + data
        prepared[path] = data
        _notify(progress, 40 + 40 * file_index / max(1, len(files)), "Preparando TL: {}/{}".format(file_index, len(files)))

    written = []
    try:
        for file_index, (path, data) in enumerate(prepared.items(), 1):
            path.write_bytes(data); written.append(path)
            _notify(progress, 80 + 15 * file_index / max(1, len(prepared)), "Gravando TL: {}/{}".format(file_index, len(prepared)))
    except Exception:
        for path in written:
            try: path.write_bytes(originals[path])
            except Exception: pass
        raise

    # The physical reconstruction map and TAB are temporary by design. They are
    # removed only after every target file has been validated and written.
    temp_dir = map_path.parent
    tab_path = temp_dir / payload.get('tab_file', 'dialogue.tab')
    if tab_path.exists(): tab_path.unlink()
    map_path.unlink()
    try: temp_dir.rmdir()
    except OSError: pass

    _log_file(output_dir, "Registros injetados: {}".format(len(entries)))
    _log_file(output_dir, "Arquivos TL alterados: {}".format(len(prepared)))
    _log_file(output_dir, "Mapa temporário removido após sucesso confirmado.")
    if log:
        log("Wells: {} traduções reinjetadas em {} arquivos TL.".format(len(entries), len(prepared)))
    _notify(progress, 100, "Reinjeção Wells concluída e mapa temporário removido.")
    return {"entries": len(entries), "files": len(prepared), "language": language, "document": str(document)}
